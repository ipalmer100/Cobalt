import pytest

from cobalt.creation import CreationError, create_blank_spec, duplicate_spec
from cobalt.docx_sections import parse_document

from .fixtures.builder import build_sample_spec_docx


def test_duplicate_spec_resets_identity_and_carries_over_data(tmp_path):
    source = str(tmp_path / "source.docx")
    build_sample_spec_docx(source, spec_number="SW0001", revision="06")

    dest = str(tmp_path / "new_spec.docx")
    new_spec = duplicate_spec(source, dest, spec_number="SW0099", customer="New Customer Inc", who="Isaac")

    assert new_spec.spec_number == "SW0099"
    assert new_spec.customer == "New Customer Inc"
    assert new_spec.revision_number == "01"

    # data tables carried over as a starting point
    bom = new_spec.primary("Bill of Materials").records()
    assert bom[0]["Raw Material"] == "48g PET"
    assert bom[0]["Supplier"] == "Flex Films"

    # revision history reset, not inherited from the source spec
    rev = new_spec.primary("Revision History").records()
    assert len(rev) == 1
    assert rev[0]["Revision #"] == "01"
    assert rev[0]["Who"] == "Isaac"
    assert "SW0001" in rev[0]["Revision"]


def test_duplicate_spec_refuses_existing_destination(tmp_path):
    source = str(tmp_path / "source.docx")
    build_sample_spec_docx(source)
    dest = tmp_path / "new_spec.docx"
    dest.write_text("already here")

    with pytest.raises(CreationError, match="already exists"):
        duplicate_spec(source, str(dest), "SW0099", "Customer", "Isaac")


def test_duplicate_spec_missing_source_raises(tmp_path):
    with pytest.raises(CreationError, match="does not exist"):
        duplicate_spec(str(tmp_path / "nope.docx"), str(tmp_path / "dest.docx"), "SW0099", "Customer", "Isaac")


def test_create_blank_spec_from_template(tmp_path):
    dest = str(tmp_path / "brand_new.docx")
    spec = create_blank_spec(dest, spec_number="SW0100", customer="Fresh Co", who="Isaac")

    assert spec.spec_number == "SW0100"
    assert spec.customer == "Fresh Co"
    assert spec.revision_number == "01"
    assert spec.warnings == []

    rev = spec.primary("Revision History").records()
    assert len(rev) == 1
    assert rev[0]["Revision"] == "Spec created from blank template."

    # blank template's data tables are genuinely empty (just headers)
    assert spec.primary("Bill of Materials").records() == []


def test_create_blank_spec_refuses_existing_destination(tmp_path):
    dest = tmp_path / "brand_new.docx"
    dest.write_text("already here")

    with pytest.raises(CreationError, match="already exists"):
        create_blank_spec(str(dest), "SW0100", "Fresh Co", "Isaac")


def test_created_specs_creates_missing_parent_directories(tmp_path):
    dest = str(tmp_path / "new_customer_folder" / "spec.docx")
    spec = create_blank_spec(dest, "SW0200", "Another Co", "Isaac")
    assert spec.spec_number == "SW0200"

    reparsed = parse_document(dest)
    assert reparsed.customer == "Another Co"


def test_a_spec_labelled_sonoco_spec_gets_its_own_number(tmp_path):
    """Specs written before the Sonoco->Toppan rename label the field
    "Sonoco Spec #". The writer only ever looked for "Spec #", found
    nothing, and reported success -- so duplicating one of those produced a
    new document still carrying the source spec's number."""
    from docx import Document

    source = tmp_path / "source.docx"
    build_sample_spec_docx(str(source), spec_number="EG1007")

    doc = Document(str(source))
    header = doc.sections[0].header.tables[0]
    for row in header.rows:
        for cell in row.cells:
            if cell.text.strip().rstrip(":").strip().lower() == "spec #":
                cell.text = "Sonoco Spec #:"
    doc.save(str(source))
    assert parse_document(str(source)).spec_number == "EG1007"

    dest = tmp_path / "new.docx"
    created = duplicate_spec(str(source), str(dest), "EG9001", "NEW CUSTOMER LTD", "Isaac")

    assert created.spec_number == "EG9001", "the new spec must not keep the source's number"
    assert created.customer == "NEW CUSTOMER LTD"
    assert parse_document(str(dest)).spec_number == "EG9001"


def test_a_toppan_labelled_spec_works_the_same_way(tmp_path):
    from docx import Document

    source = tmp_path / "source.docx"
    build_sample_spec_docx(str(source), spec_number="EG1008")
    doc = Document(str(source))
    for row in doc.sections[0].header.tables[0].rows:
        for cell in row.cells:
            if cell.text.strip().rstrip(":").strip().lower() == "spec #":
                cell.text = "Toppan Spec #:"
    doc.save(str(source))

    dest = tmp_path / "new.docx"
    assert duplicate_spec(str(source), str(dest), "EG9002", "ACME", "Isaac").spec_number == "EG9002"


def test_creation_refuses_rather_than_produce_a_spec_with_the_wrong_number(tmp_path):
    """No spec-number field at all: better to refuse than to hand someone a
    document that claims to be the spec it was copied from."""
    from docx import Document

    source = tmp_path / "source.docx"
    build_sample_spec_docx(str(source), spec_number="EG1009")
    doc = Document(str(source))
    for row in doc.sections[0].header.tables[0].rows:
        for cell in row.cells:
            if cell.text.strip().rstrip(":").strip().lower() == "spec #":
                cell.text = "Reference:"
    doc.save(str(source))

    dest = tmp_path / "new.docx"
    with pytest.raises(CreationError, match="spec-number field"):
        duplicate_spec(str(source), str(dest), "EG9003", "ACME", "Isaac")
